# -*- coding: utf-8 -*-
"""三端点名门的脚手架：同一个数在图纸、方案书、读我里各印了一遍，这里现算比对，
并且自带破坏测试——按一遍每一条判据，故意把它那一端的原文抠掉／把数字改一位，
门必须响；不响的门就是瞎门，交付看着全绿其实没人守着。

为什么要有这一支：交付里的数字是抄出来的，抄就会走样。本项目真出过的事：方案书写
"保留 24 株"、图纸图签写"保留 23 株"、读我又写回 24——肉眼翻 19 张图 60 页文档发现不了，
机器一条正则 3 秒抓出来。更坑的是判据自己写错了（正则永远匹配不到、期望值恒等），
表现和"全对"一模一样，所以要靠破坏测试证明门是活的。

用法：
    python gates_scaffold.py --demo                       # 内置示例跑一遍（含自测）
    python gates_scaffold.py --spec gates.json --textdir texts/          # 正式比对
    python gates_scaffold.py --spec gates.json --textdir texts/ --selftest
    python gates_scaffold.py --fresh 产物表.csv --textdir texts/         # 顺带查时效性

texts/ 里每个文件是一"端"：文件名去掉扩展名就是端名（图纸.dxf → 端「图纸」）。
.dxf 取 modelspace 全部 TEXT，.docx 取正文与表格文字，.txt/.md 原样。同一端多个文件
同名不同扩展会被合并。

gates.json 形如：
    [{"name": "保留树株数", "pats": {"图纸": "(?P<n>\\\\d+)株保留", "正文": "保留(?P<n>\\\\d+)株"},
      "exp": {"n": 24}}]
exp 里的期望值不要手抄，用现算的表达式填（Python 里 f-string 或 json.dump 前算好）；
手抄的期望值本身就是第二个会走样的地方。
"""
import argparse
import glob
import json
import os
import re
import sys
from pathlib import Path


def tight(s):
    """抹掉一切空白与全角空格：被匹配的文本是拼接后压平的，判据里留一个空格就永远
       匹配不到，而报出来的是「判据缺失」——看着像哪一端漏写了，其实是正则自己带的空格。"""
    return re.sub(r"\s+", "", str(s).replace("　", ""))


def same(v, e):
    """容差由印出来的小数位数决定：印 "12.84" 只允许 ±0.005，印 "153" 允许 ±0.5。
       宽容差会把"苗木占地 12.84 印成 12.3"也放过去，所以不写死容差表。"""
    if isinstance(e, str):
        return e == v
    try:
        f = float(v)
    except (TypeError, ValueError):
        return str(e) == v
    d = len(v.split(".")[1]) if "." in v else 0
    return abs(f - float(e)) <= 0.5 * 10.0 ** (-d) + 1e-9


def load_ends(textdir):
    """一"端"一份文本。DXF／DOCX 都只取文字，端与端可比。"""
    out = {}
    for p in sorted(Path(textdir).iterdir()):
        if p.is_dir():
            continue
        end, suf = p.stem, p.suffix.lower()
        if suf == ".dxf":
            import ezdxf
            d = ezdxf.readfile(str(p))
            parts = []
            for e in d.modelspace().query("TEXT MTEXT ATTDEF"):
                # MTEXT 的 dxf.text 带 {\C1;…} 这类格式码，直接拼进文本会让判据匹配不到；
                # plain_text() 才是图上真正显示的那串字。
                parts.append(e.plain_text() if hasattr(e, "plain_text") else e.dxf.text)
            txt = " ".join(parts)
        elif suf == ".docx":
            import docx
            doc = docx.Document(str(p))
            txt = " ".join(x.text for x in doc.paragraphs) + " " + " ".join(
                c.text for t in doc.tables for r in t.rows for c in r.cells)
        elif suf in (".txt", ".md", ".json"):
            txt = p.read_text(encoding="utf-8", errors="replace")
        else:
            continue
        out[end] = tight(out.get(end, "") + txt)
    assert out, "%s 里没有任何可读的一「端」（支持 .dxf/.docx/.txt/.md/.json）" % textdir
    return out


def match_all(C, T):
    """返回（红字列表，判据条数）。跑正式比对和跑自测用的是同一支笔——自测里验过的
       就是正式跑的那条，不会出现「自测绿、正式瞎」。"""
    bad, n = [], 0
    for g in C:
        for end, pat in g["pats"].items():
            n += 1
            if end not in T:
                bad.append("判据缺失｜%s｜%s 这一端根本没加载到" % (g["name"], end))
                continue
            m = re.search(tight(pat), T[end])
            if not m:
                bad.append("判据缺失｜%s｜%s" % (g["name"], end))
                continue
            for k, v in (m.groupdict() or {}).items():
                if k in g.get("exp", {}) and not same(v, g["exp"][k]):
                    bad.append("三端不符｜%s｜%s 印 %s，现算 %s"
                               % (g["name"], end, v, g["exp"][k]))
    return bad, n


# ── 破坏测试：每条判据都要能被自己弄响 ────────────────────────────────────────────
def drop(T, g, end):
    """把这条判据在这一端匹配到的原文整段抠掉：门是死的就不会报缺失。
       抠全部匹配而不是第一处——同一句话在正文出现两回，只删一处另一处还在。"""
    pat = tight(g["pats"][end])
    lit = [m.group(0) for m in re.finditer(pat, T[end])]
    assert lit, "破坏无处可下：%s｜%s 本来就没匹配到" % (g["name"], end)
    t = T[end]
    for s in lit:
        t = t.replace(s, "")
    return dict(T, **{end: t})


def strip_head(T, g, end):
    """链式判据（A.*?B.*?C）专用的破坏动作。drop() 把 pattern 当字面去抠文本，长链抠不动，
       于是自测报"破坏无处可下"——看着像门坏了，其实是动作选错了。这里只抠掉第一个通配符
       之前的那段头部：头部一没，整条链在这一端必然断，且不会误伤别处的同字样。"""
    pat = tight(g["pats"][end])
    m = re.search(pat, T[end])
    assert m, "破坏无处可下：%s｜%s 本来就没匹配到" % (g["name"], end)
    cut = re.search(r"\.\*\?|\.\+", pat)
    head = pat[:cut.start()] if cut else pat
    assert head and head in T[end], "头部不是字面（%s），换 drop() 或 bump()" % g["name"]
    return dict(T, **{end: T[end].replace(head, "", 1)})


def bump(T, g, end):
    """把这一端印出的第一个数加一个最小位（整数 +1、两位小数 +0.01）——用判据自己的
       正则定位，不另抄一遍文案。"""
    pat = tight(g["pats"][end])
    m = re.search(pat, T[end])
    assert m, "自测找不到这条判据的原文：%s｜%s" % (g["name"], end)
    ks = [k for k in (m.groupdict() or {})
          if re.fullmatch(r"\d+(\.\d+)?", m.group(k) or "") and k in g.get("exp", {})]
    assert ks, "这条判据没有可比对的数字，破坏不了（%s）" % g["name"]
    k = ks[0]
    s, i = m.group(k), m.start(k)
    d = len(s.split(".")[1]) if "." in s else 0
    return dict(T, **{end: T[end][:i] + "%.*f" % (d, float(s) + 10.0 ** -d) + T[end][i + len(s):]})


def selftest(C, T):
    """→ (响过的门数, 门总数, 瞎门列表)。判据名要能在破坏后的红字里出现，才算活着。"""
    dead, alive = [], 0
    for g in C:
        fired = False
        for end in g["pats"]:
            if end not in T:
                continue
            for fn in (drop, bump, strip_head):
                try:
                    bad, _ = match_all([g], fn(T, g, end))
                except AssertionError:
                    continue
                if any(g["name"] in q for q in bad):
                    fired = True
        if fired:
            alive += 1
        else:
            dead.append("%s（pats=%s）" % (g["name"], list(g["pats"])))
    return alive, len(C), dead


def fresh(mapcsv):
    """时效性：改过上游脚本必须重跑那一环节。产物比脚本旧，就是产物是旧的——
       这类红字最容易被当成「跑一次门太麻烦」忽略，而它恰好意味着你交付的是上一版的结果。"""
    import csv
    bad, n = [], 0
    with open(mapcsv, newline="", encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            n += 1
            prod, src = r.get("产物", "").strip(), r.get("生成脚本", "").strip()
            if not (prod and src):
                bad.append("时效性表缺列｜产物／生成脚本（第 %d 行）" % n)
                continue
            if not os.path.isfile(prod):
                bad.append("产物不存在｜%s" % prod)
                continue
            if not os.path.isfile(src):
                bad.append("生成脚本不存在｜%s（%s 声称由它生成）" % (src, prod))
                continue
            if os.path.getmtime(src) > os.path.getmtime(prod) + 1e-6:
                bad.append("产物过期｜%s 比 %s 旧，重跑那一环节" % (prod, src))
    return n, bad


AREA_M2 = 5347.0            # 现场量的红线面积（数据层里的真值）

DEMO_C = [
    {"name": "保留树株数",
     "pats": {"图纸": r"(?P<n>\d+)株保留", "正文": r"保留(?P<n>\d+)株", "读我": r"保留(?P<n>\d+)株"},
     "exp": {"n": 24}},
    {"name": "常绿比例",
     "pats": {"图纸": r"常绿(?P<p>\d+)%", "正文": r"常绿(?P<p>\d+)%"},
     "exp": {"p": 70}},
    {"name": "红线面积到公顷",
     "pats": {"正文": r"红线面积(?P<a>\d+\.\d\d)公顷"},
     # 期望值现算，不手抄：手抄的期望值本身就是第二个会走样的地方
     "exp": {"a": round(AREA_M2 / 10000.0, 2)}},
]


def demo():
    T = {"图纸": tight("种植设计总平面图 1:500 红线面积0.53公顷 24株保留 常绿70%"),
         "正文": tight("本方案红线面积0.53公顷，保留24株现状乔木，配比常绿70%落叶30%。"),
         "读我": tight("图集19张，保留24株，详见PL-06苗木表。")}
    print("内置示例：3 条判据、%d 端文本（数据层里 AREA_M2=%.0f 现算出 0.53 公顷）"
          % (len(T), AREA_M2))
    bad, n = match_all(DEMO_C, T)
    print("正式比对：%d 个判据点，%d 条红字%s" % (n, len(bad), "（先确认全绿，再看下面的破坏）"))
    for q in bad:
        print("  ✗ " + q)
    alive, tot, dead = selftest(DEMO_C, T)
    print("破坏测试：%d/%d 道门响过%s" % (alive, tot, "" if not dead else "，瞎门：" + "、".join(dead)))
    print("\n现在把「图纸」端的 24 株改成 23，看门响不响：")
    T2 = dict(T, 图纸=T["图纸"].replace("24株保留", "23株保留"))
    for q in match_all(DEMO_C, T2)[0]:
        print("  ✗ " + q)
    print("\n再把「正文」端那句「保留24株」整个删掉：")
    T3 = dict(T, 正文=T["正文"].replace("保留24株", ""))
    for q in match_all(DEMO_C, T3)[0]:
        print("  ✗ " + q)
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--demo", action="store_true", help="跑内置示例")
    ap.add_argument("--spec", help="判据表 gates.json")
    ap.add_argument("--textdir", help="各端文本目录")
    ap.add_argument("--fresh", help="时效性表 CSV：产物,生成脚本")
    ap.add_argument("--selftest", action="store_true", help="对每条判据做破坏测试")
    ap.add_argument("--json", metavar="OUT")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    if a.demo or not a.spec:
        return demo()
    C = json.loads(Path(a.spec).read_text(encoding="utf-8"))
    T = load_ends(a.textdir) if a.textdir else {}
    bad, n = match_all(C, T)
    rep = {"判据点": n, "端": sorted(T), "红字": bad}
    print("点名 %d 条判据／%d 个判据点，红字 %d 条" % (len(C), n, len(bad)))
    for q in bad:
        print("  ✗ " + q)
    if a.selftest:
        alive, tot, dead = selftest(C, T)
        rep["门不瞎"] = "%d/%d" % (alive, tot)
        print("破坏测试：%d/%d 道门响过" % (alive, tot))
        for d in dead:
            print("  ✗ 瞎门｜%s" % d)
        bad += ["瞎门｜" + d for d in dead]
    if a.fresh:
        fn, fb = fresh(a.fresh)
        rep["时效"] = "%d/%d" % (fn - len(fb), fn)
        print("时效性：%d/%d 条新鲜" % (fn - len(fb), fn))
        for q in fb:
            print("  ✗ " + q)
        bad += fb
    if a.json:
        Path(a.json).write_text(json.dumps(rep, ensure_ascii=False, indent=1), encoding="utf-8")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
