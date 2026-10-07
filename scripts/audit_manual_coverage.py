# -*- coding: utf-8 -*-
"""现读《指导手册／任务书》原文，把手册自己的图目抓出来，和你的图纸清单对账。

为什么要现读而不是抄一份清单：手册是作业要求的唯一真源，而且它会变（上一版 图4 只有
一张，这一版拆成 图4-1／图4-2）。抄清单等于把「齐不齐」这句话交给一段会过期的文字。
这道检查在交付前跑，红字为 0 才可以说"手册要求的图全齐"。

对照表 CSV 至少两列（带表头，允许 BOM，多余列忽略）：
    手册图号,承载图纸,承载文件            承载文件是相对 --root 的路径，图号形如 图4-1
可选列 手册图名：给了就会去手册原文里核这三个字在不在（防止你把"没要求的图"当成已完成）。

用法：
    python audit_manual_coverage.py --manual 指导手册.docx --manifest 对照表.csv --root .
    python audit_manual_coverage.py --manual 指导手册.pdf  --manifest 对照表.csv --json out.json
"""
import argparse
import csv
import json
import re
import sys
from pathlib import Path

CAP = re.compile(r"^(图\s*\d+(?:[-－]\d+)?)")      # 行首是"图N"才算图注
BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"


def tight(s):
    return re.sub(r"\s+", "", str(s).replace("　", ""))


def _nimg(el):
    return len(el.findall(".//" + BLIP))


def _docx_rows(path):
    """正文段与表格行拉平成一条序列，每行给（文本, 这一行挂了几张图）。"""
    import docx
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = docx.Document(str(path))
    for ch in doc.element.body.iterchildren():
        if isinstance(ch, CT_P):
            yield Paragraph(ch, doc).text.strip(), _nimg(ch)
        elif isinstance(ch, CT_Tbl):
            for row in Table(ch, doc).rows:
                cell = row.cells[0]
                txt = "|".join(p.text.strip() for p in cell.paragraphs).strip()
                yield txt, _nimg(cell._tc)


def _pdf_rows(path):
    import pypdfium2 as pdfium
    d = pdfium.PdfDocument(str(path))
    for i in range(len(d)):
        for line in d[i].get_textpage().get_text_range().splitlines():
            yield line.strip(), 0            # PDF 认不出图挂在哪个编号上，只认编号


def manual_figures(path):
    """→ [(图号, 配到的图片数, 图注原文), …]，按手册里出现的先后排。
       配对规则：图注行上方尚未认领的图片归它——两处版式都是"图在上、号在下"。"""
    src = _docx_rows(path) if str(path).lower().endswith(".docx") else _pdf_rows(path)
    out, pend = [], 0
    for txt, n in src:
        m = CAP.match(txt)
        if m:
            out.append((m.group(1).replace(" ", ""), pend, txt))
            pend = 0
        elif n:
            pend += n
    return out


def manual_fulltext(path):
    if str(path).lower().endswith(".docx"):
        return tight(" ".join(t for t, _n in _docx_rows(path)))
    import pypdfium2 as pdfium
    d = pdfium.PdfDocument(str(path))
    return tight(" ".join(d[i].get_textpage().get_text_range() for i in range(len(d))))


def load_manifest(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = [{(k or "").strip(): (v or "").strip() for k, v in r.items()}
                for r in csv.DictReader(f)]
    for i, r in enumerate(rows):
        if "手册图号" not in r or "承载文件" not in r:
            raise SystemExit("对照表第 %d 行起缺列：表头必须有 手册图号／承载文件（%s）"
                             % (i + 1, sorted(r)))
    return rows


def audit(manual, manifest, root):
    figs = manual_figures(manual)
    labs = {}
    for lab, n, txt in figs:
        labs[lab] = max(n, labs.get(lab, 0))
    rows = load_manifest(manifest)
    full = None
    hard, warn = [], []

    covered = {r["手册图号"].replace(" ", "") for r in rows}
    for lab in sorted(labs, key=lambda s: [int(x) for x in re.split(r"[-－]", s.replace("图", ""))]):
        if lab not in covered:
            hard.append("手册要求｜%s 在手册里有图目，对照表里没有承载文件" % lab)
    if any("手册图名" in r and r["手册图名"] for r in rows):
        full = manual_fulltext(manual)
    for r in rows:
        lab = r["手册图号"].replace(" ", "")
        if lab not in labs:
            warn.append("对照表多列｜%s 手册里根本没有这个图号（%s）——是自己加的图就写清出处"
                        % (lab, r.get("承载图纸", "")))
        p = (root / r["承载文件"]).resolve()
        if not p.is_file():
            hard.append("承载文件不存在｜%s → %s" % (lab, r["承载文件"]))
        nm = r.get("手册图名", "")
        if nm and full and tight(nm) not in full:
            hard.append("手册图名对不上｜对照表写「%s」，手册原文里没有这三个字（%s）" % (nm, lab))
        if labs.get(lab, -1) == 0:
            warn.append("手册这一行没配图｜%s（%s）——核对是不是正文里的引用而非图目"
                        % (lab, r.get("承载图纸", "")))
    from collections import Counter
    c = Counter(tight(r["承载文件"]) for r in rows if r["承载文件"])
    for f, n in sorted(c.items()):
        if n > 1:
            warn.append("一文件认领 %d 个手册图号｜%s——确认这张图真的画了这几块内容" % (n, f))
    return {"手册图目": [{"图号": l, "图片数": n, "图注": t} for l, n, t in figs],
            "对照表行": len(rows), "硬伤": hard, "提醒": warn}


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--manual", required=True, help="指导手册／任务书 .docx 或 .pdf")
    ap.add_argument("--manifest", required=True, help="图纸对照表 .csv")
    ap.add_argument("--root", default=".", help="承载文件的相对根目录")
    ap.add_argument("--json", metavar="OUT", help="同时写一份 JSON")
    a = ap.parse_args()
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    r = audit(Path(a.manual), Path(a.manifest), Path(a.root))
    if a.json:
        Path(a.json).write_text(json.dumps(r, ensure_ascii=False, indent=1), encoding="utf-8")
    print("手册图目 %d 条／对照表 %d 行" % (len(r["手册图目"]), r["对照表行"]))
    for x in r["手册图目"]:
        print("  %-7s 配图 %d  %s" % (x["图号"], x["图片数"], x["图注"][:38]))
    for tag, key in (("硬伤", "硬伤"), ("提醒", "提醒")):
        print("\n%s %d 条" % (tag, len(r[key])))
        for q in r[key]:
            print("  ✗ " + q if tag == "硬伤" else "  · " + q)
    return 1 if r["硬伤"] else 0


if __name__ == "__main__":
    sys.exit(main())
